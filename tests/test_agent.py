from datetime import datetime, timezone, timedelta
import json
import zipfile
import pytest
from docx import Document
from jobagent.core import Job, Store, freshness, canonical, Analysis
from jobagent.tailor import facts, validate, render
from jobagent.digest import safe_cell, build
from jobagent.apply import prefill, allowed_host

AT = datetime(2026, 10, 9, 3, 30, tzinfo=timezone.utc)


def job(**kw):
    data = dict(title='Site Reliability Engineer', company='Example', location='India',
                url='https://jobs.lever.co/example/123', source='fixture',
                description='Linux production reliability engineering',
                posted_at=AT.isoformat(), time_kind='published_exact')
    return Job(**(data | kw))


@pytest.mark.parametrize('hours,expected', [(0, 'verified'), (12, 'verified'), (12.01, 'outside_window'), (-1, 'outside_window')])
def test_window(hours, expected):
    assert freshness(job(posted_at=(AT-timedelta(hours=hours)).isoformat()), AT) == expected


def test_unverified():
    for kind in ['date_only', 'updated', 'unknown']:
        assert freshness(job(time_kind=kind), AT) == 'unverified'
    assert freshness(job(posted_at='2026-10-09'), AT) == 'unverified'


def test_canonical():
    assert canonical('https://in.indeed.com/viewjob?utm_source=x&jk=123') == 'https://in.indeed.com/viewjob?jk=123'
    assert job().id == job(url=job().url+'?utm_source=elsewhere').id
    with pytest.raises(ValueError): canonical('file:///etc/passwd')


def test_dedupe_and_state(tmp_path):
    db = Store(tmp_path / 'jobs.db')
    j = job()
    db.add(j, AT)
    db.update(j.id, status='applied', applied_at=AT.isoformat())
    db.add(j, AT)
    assert len(db.rows()) == 1
    assert db.get(j.id)['status'] == 'applied'
    assert Store(tmp_path / 'jobs.db').get(j.id)['applied_at']


def test_unknown_evidence_rejected():
    a = Analysis(score=90, location_eligible=True, location_reason='India', missing_skills=[], rationale='fit', highlight_ids=[999])
    with pytest.raises(ValueError): validate(a, {5: 'real fact'})


def test_resume_preserves_claims_and_employer_boundaries(tmp_path):
    d = Document()
    for s in ['Candidate', 'Engineer', 'contact', '', 'EXPERIENCE', 'Employer A']:
        d.add_paragraph(s)
    first = 'Monitored Linux production services and responded to incidents with the support team every week.'
    second = 'Built observability dashboards using existing metrics to assist with production incident investigation.'
    d.add_paragraph(first).runs[0].bold = True
    d.add_paragraph(second)
    d.add_paragraph('Employer B')
    d.add_paragraph('Operated network infrastructure and performed incident response within established change controls.')
    src, dst = tmp_path/'source.docx', tmp_path/'tailored.docx'
    d.save(src)
    a = Analysis(score=80, location_eligible=True, location_reason='India', missing_skills=['Kubernetes'], rationale='fit', highlight_ids=[7])
    render(src, dst, a, facts(src))
    out = Document(dst)
    assert sorted(p.text for p in out.paragraphs) == sorted(p.text for p in d.paragraphs)
    assert out.paragraphs[6].text == second
    assert out.paragraphs[8].text == 'Employer B'
    assert out.paragraphs[7].runs[0].bold


def test_spreadsheet_injection():
    assert safe_cell('=HYPERLINK("evil")').startswith("'")


def test_daily_applied_report_uses_ist(tmp_path):
    db = Store(tmp_path/'db')
    j = job()
    db.add(j, AT-timedelta(days=1))
    db.update(j.id, status='applied', applied_at=AT.isoformat(), evidence='confirmation')
    _, files, _ = build(db, tmp_path/'out', AT, [])
    from openpyxl import load_workbook
    rows = list(load_workbook(files[0]).active.values)
    assert rows[1][4] == 'applied'
    assert '2026-10-09' in files[0].name


def test_prefill_has_no_clicks_or_submit():
    class Field:
        def count(self): return 1
        def is_visible(self): return True
        def is_editable(self): return True
        def fill(self, value): self.value = value
    class Page:
        def __init__(self): self.field = Field()
        def get_by_label(self, pattern): return self.field
    page = Page()
    assert prefill(page, {'email': 'test@example.com'}) == (['email'], [])
    assert page.field.value == 'test@example.com'
    assert not allowed_host('lever.co.attacker.example')


def test_offline_run_no_mail_no_gemini(tmp_path, monkeypatch):
    from jobagent.cli import run
    import jobagent.tailor
    import jobagent.digest
    def forbidden(*a, **k): raise AssertionError('external call in dry run')
    monkeypatch.setattr(jobagent.tailor, 'analyze', forbidden)
    monkeypatch.setattr(jobagent.digest, 'send', forbidden)
    d = Document(); d.add_paragraph('Candidate'); d.save(tmp_path/'resume.docx')
    fixture = tmp_path/'jobs.json'
    from jobagent.core import now
    fixture.write_text(json.dumps([job(posted_at=now().isoformat()).model_dump()]))
    cfg = {'resume': 'resume.docx', 'include_unverified': False, 'max_jobs_per_run': 8}
    run(cfg, tmp_path, True, str(fixture))
    assert not (tmp_path/'state/jobs.db').exists()
    assert (tmp_path/'dry-run/output/digest.txt').exists()


def test_historical_application_is_not_lost_at_midnight(tmp_path):
    from openpyxl import load_workbook
    db = Store(tmp_path/'db')
    j = job()
    db.add(j, AT-timedelta(days=2))
    db.update(j.id, status='applied', applied_at=(AT-timedelta(days=1)).isoformat(), evidence='receipt')
    _, files, _ = build(db, tmp_path/'out', AT, [])
    wb = load_workbook(files[0])
    assert list(wb['Applied history'].values)[1][0] == j.id


def test_source_parsers_use_publication_not_update(monkeypatch):
    import jobagent.sources as sources
    def get(url, **kw):
        if url.endswith('/jobs'):
            return {'jobs': [{'id': 123, 'title': 'DevOps Engineer', 'location': {'name': 'India'},
                             'absolute_url': 'https://boards.greenhouse.io/example/jobs/123',
                             'content': '<p>Linux</p>', 'updated_at': AT.isoformat()}]}
        return {'first_published': (AT-timedelta(days=5)).isoformat()}
    monkeypatch.setattr(sources, 'get', get)
    jobs, errors = sources.discover({'sites': [], 'roles': [], 'locations': [], 'greenhouse': ['example']})
    assert not errors
    assert freshness(jobs[0], AT) == 'outside_window'


def test_source_failure_is_reported(monkeypatch):
    import jobagent.sources as sources
    def fail(*a, **k): raise TimeoutError('remote failure')
    monkeypatch.setattr(sources, 'get', fail)
    jobs, errors = sources.discover({'sites': [], 'roles': [], 'locations': [], 'lever': ['example']})
    assert jobs == []
    assert errors == ['lever/example: TimeoutError']
