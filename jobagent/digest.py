import json
import os
import smtplib
import ssl
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from .core import Job


def safe_cell(value):
    # Prevent spreadsheet formula injection from externally supplied job text.
    s = str(value or '')
    return "'" + s if s.startswith(('=', '+', '-', '@')) else s


def build(store, out, at, errors):
    out.mkdir(parents=True, exist_ok=True)
    today = at.astimezone(ZoneInfo('Asia/Kolkata')).date()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Daily activity'
    ws.append(['Job ID', 'Title', 'Company', 'Location', 'Status', 'Publication confidence', 'Application URL', 'Applied at', 'Submission evidence'])
    pending = []
    for row in store.rows():
        job = Job.model_validate_json(row['payload'])
        seen_today = datetime.fromisoformat(row['first_seen']).astimezone(ZoneInfo('Asia/Kolkata')).date() == today
        applied_today = row['applied_at'] and datetime.fromisoformat(row['applied_at']).astimezone(ZoneInfo('Asia/Kolkata')).date() == today
        if seen_today or applied_today:
            ws.append([safe_cell(v) for v in [row['id'], job.title, job.company, job.location, row['status'], row['freshness'], job.url, row['applied_at'], row['evidence']]])
        if row['status'] == 'prepared' and not row['emailed_at']:
            pending.append(row)
    history = wb.create_sheet('Applied history')
    history.append(['Job ID', 'Title', 'Company', 'Application URL', 'Applied at (UTC)', 'Evidence'])
    for row in store.rows():
        if row['status'] == 'applied':
            j = Job.model_validate_json(row['payload'])
            history.append([safe_cell(v) for v in [row['id'], j.title, j.company, j.url, row['applied_at'], row['evidence']]])
    history.freeze_panes = 'A2'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(65, max(16, max(len(str(c.value or '')) for c in col) + 2))
    health = wb.create_sheet('Source health')
    health.append(['Source errors (a failure is not zero matches)'])
    for e in errors:
        health.append([safe_cell(e)])
    report = out / f'job-activity-{today}.xlsx'
    wb.save(report)
    lines = [f'Job agent — {today} IST', 'Prepared resumes require your review. Applied means manually confirmed submission.', '']
    files = [report]
    for row in pending:
        job = Job.model_validate_json(row['payload'])
        a = json.loads(row['analysis'])
        lines.extend([f'{row["id"]}: {job.title} | {job.company} | fit {a["score"]}/100', job.url,
                      f'Publication: {row["freshness"]}', 'Missing skills: ' + ', '.join(a['missing_skills']),
                      f'Open helper: job-agent apply {row["id"]}', ''])
        files.append(Path(row['resume']))
    unverified = [r for r in store.rows() if r['freshness'] == 'unverified' and r['status'] == 'time_unverified']
    lines.append(f'{len(unverified)} postings held for unverified publication times. See report / job-agent list.')
    rating_pending = sum(r['status'] == 'rating_pending' for r in store.rows())
    lines.append(f'{rating_pending} companies awaiting verified rating evidence (must be above 3.5/5).')
    if errors:
        lines.extend(['SOURCE ERRORS:', *errors])
    text = '\n'.join(lines)
    (out / 'digest.txt').write_text(text)
    return text, files, pending


def send(text, files):
    sender = os.environ['GMAIL_ADDRESS']
    recipient = os.environ['DIGEST_TO']
    if not sender or not recipient:
        raise ValueError('Gmail sender and recipient must be configured')
    msg = EmailMessage()
    msg['From'], msg['To'], msg['Subject'] = sender, recipient, 'Job matches, reviewed resumes and daily activity'
    msg.set_content(text)
    if sum(p.stat().st_size for p in files) > 17_000_000:
        raise ValueError('Attachments too large for safe Gmail delivery; reduce max_jobs_per_run')
    for path in files:
        subtype = 'vnd.openxmlformats-officedocument.wordprocessingml.document' if path.suffix == '.docx' else 'vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        msg.add_attachment(path.read_bytes(), maintype='application', subtype=subtype, filename=path.name)
    with smtplib.SMTP_SSL('smtp.gmail.com', 465, context=ssl.create_default_context(), timeout=40) as smtp:
        smtp.login(sender, os.environ['GMAIL_APP_PASSWORD'])
        smtp.send_message(msg)
